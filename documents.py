import base64
import io
import zipfile
from pathlib import Path

def extract(filename, encoded):
    data = base64.b64decode(encoded, validate=True)
    if len(data) > 10 * 1024 * 1024:
        raise ValueError('文件不能超过 10MB')
    suffix = Path(filename).suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError('请上传未加密 PDF')
        if len(reader.pages) > 100:
            raise ValueError('最多支持 100 页')
        text = '\n'.join(p.extract_text() or '' for p in reader.pages)
    elif suffix == '.docx':
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError('解压后文件过大')
        from docx import Document
        doc = Document(io.BytesIO(data))
        text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
    else:
        raise ValueError('仅支持 PDF / DOCX')
    if not text.strip():
        raise ValueError('未提取到文字，扫描 PDF 请先 OCR 或粘贴文本')
    return text[:100000]
