"""Source-derived Word generation. Never copy sample candidate facts or photos."""
import copy
import io
from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.ns import qn
from profile_v2 import PERSONAL,EDUCATION,EXPERIENCE

def text_of(el):return ''.join(el.itertext()) if False else ''.join(t.text or '' for t in el.iter(qn('w:t')))

def replace_paragraph(p,text):
    sample=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for child in list(p._p):
        if child.tag!=qn('w:pPr'):p._p.remove(child)
    run=p.add_run(text)
    if sample is not None:run._r.insert(0,sample)

def cell_text(cell,text):
    p=cell.paragraphs[0]
    for other in list(cell.paragraphs)[1:]:other._p.getparent().remove(other._p)
    replace_paragraph(p,text)

def template_parts(doc):
    if len(doc.tables)!=3 or len(doc.tables[0].columns)!=2 or len(doc.tables[1].columns)!=4 or len(doc.tables[2].columns)!=3:
        raise ValueError('模板需要与阿迪门店模板一致的三张表：个人信息、教育、经历概况。')
    labels=[r.cells[0].text.strip().rstrip(':：').replace(' ','') for r in doc.tables[0].rows]
    if labels!=['年龄','性别','城市','出生年月','手机号','邮箱','身高体重','婚育','当前薪资','期望薪资']:
        raise ValueError('个人信息字段顺序与已确认模板不同，请使用阿迪门店结构的模板。')
    paras=doc.paragraphs
    first=next(i for i,p in enumerate(paras) if p.text.strip().startswith('工作经历'))
    detail=next(i for i,p in enumerate(paras) if i>first and p.text.strip().startswith('工作经历'))
    duty=next(i for i,p in enumerate(paras) if i>detail and p.text.strip().startswith('工作职责'))
    performance=next(i for i,p in enumerate(paras) if i>duty and p.text.strip().startswith('业绩'))
    return detail,duty,performance

def render_profile(template,profile):
    doc=Document(template);detail,duty,performance=template_parts(doc)
    p=profile['personal'];paras=doc.paragraphs
    # Store formatting patterns before replacing the repeated block.
    patterns=[copy.deepcopy(paras[i]._p) for i in [detail+1,detail+2,duty,duty+1,performance,performance+1]]
    # Template title is retained; all content between title and personal table is candidate name/spacing.
    first_table=doc.tables[0]._tbl
    seen_title=False;name_done=False
    for el in list(doc.element.body):
        if el is first_table:break
        if el.tag==qn('w:p'):
            para=Paragraph(el,doc)
            if not seen_title and para.text.strip():seen_title=True;continue
            if para.text.strip() and not name_done:replace_paragraph(para,p.get('姓名',''));name_done=True
            else:replace_paragraph(para,'')
    values={k:p.get(k,'') for k in PERSONAL};values['身高体重']=' / '.join(v for v in [p.get('身高',''),p.get('体重','')] if v);values['期望薪资']=''
    for row in doc.tables[0].rows:
        label=row.cells[0].text.strip().rstrip(':：').replace(' ','')
        cell_text(row.cells[1],values.get(label,''))
    def rows(table,records,keys,header=False):
        sample=copy.deepcopy(table.rows[1 if header else 0]._tr)
        for row in list(table.rows)[1 if header else 0:]:table._tbl.remove(row._tr)
        for obj in records or [dict.fromkeys(keys,'')]:
            table._tbl.append(copy.deepcopy(sample));row=table.rows[-1]
            for cell,key in zip(row.cells,keys):cell_text(cell,obj.get(key,''))
            # Prevent fixed-height truncation inherited from input templates.
            if row._tr.trPr is not None:
                for h in list(row._tr.trPr.findall(qn('w:trHeight'))):row._tr.trPr.remove(h)
    rows(doc.tables[1],profile['education'],EDUCATION)
    rows(doc.tables[2],profile['experience'],['任职时间','公司','职位'],True)
    anchor=paras[detail]._p;found=False
    for el in list(doc.element.body):
        if el is anchor:found=True;continue
        if found and el.tag!=qn('w:sectPr'):doc.element.body.remove(el)
    def append(pattern,text,keep=False):
        nonlocal anchor
        el=copy.deepcopy(pattern);anchor.addnext(el);anchor=el
        para=Paragraph(el,doc);replace_paragraph(para,text)
        para.paragraph_format.keep_with_next=keep
        para.paragraph_format.keep_together=False
    for e in profile['experience'] or [dict.fromkeys(EXPERIENCE,'')]:
        block_start=anchor
        append(patterns[0],e.get('任职时间','')+'\t'+e.get('公司',''),True)
        append(patterns[1],e.get('职位',''),True)
        append(patterns[2],'工作职责：',True);append(patterns[3],e.get('工作职责',''))
        if e.get('拥有技能'):append(patterns[3],'拥有技能：'+e['拥有技能'])
        append(patterns[4],'业绩：',True);append(patterns[5],e.get('业绩',''))
        for store in e.get('stores',[]):
            content='；'.join(k+'：'+v for k,v in store.items() if v)
            if content:append(patterns[3],content)
        append(patterns[3],'离职原因：________________')
        # Keep short experience blocks together so a lone leave-reason line cannot spill onto a new page.
        if len(str(e))<1100:
            node=block_start.getnext()
            while node is not None and node is not anchor:
                Paragraph(node,doc).paragraph_format.keep_with_next=True
                node=node.getnext()
    append(patterns[2],'主观评价：',True);append(patterns[3],'________________')
    # Remove body photos / drawings and their image parts, but preserve header/footer branding.
    for el in list(doc.element.body.iter()):
        if el.tag in [qn('w:drawing'),qn('w:pict')]:el.getparent().remove(el)
    used={v for el in doc.element.iter() for k,v in el.attrib.items() if k.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}')}
    for rid,rel in list(doc.part.rels.items()):
        if rel.reltype.endswith('/customXml') or (rel.reltype.endswith('/image') and rid not in used):doc.part.drop_rel(rid)
    for key in ['author','last_modified_by','comments','keywords','subject','title']:setattr(doc.core_properties,key,'')
    output=io.BytesIO();doc.save(output);return output.getvalue()

def sanitize_template(source,dest):
    from profile_v2 import empty_profile
    p=empty_profile();p['personal']['姓名']='候选人姓名'
    p['education']=[dict.fromkeys(EDUCATION,'')]
    p['experience']=[dict(dict.fromkeys(EXPERIENCE,''),stores=[])]
    Path(dest).parent.mkdir(parents=True,exist_ok=True)
    Path(dest).write_bytes(render_profile(source,p))
