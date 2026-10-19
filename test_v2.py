import unittest,tempfile,io,base64,sqlite3,json,zipfile
from pathlib import Path
from datetime import datetime,timedelta,timezone
from docx import Document
from unittest.mock import patch
from word_v2 import render_profile
from store_v2 import V2Store
from profile_v2 import local_extract,missing

RESUME='''姓名：测试候选人A
城市：上海
教育背景：
2014-2018 | 虚构大学 | 管理学 | 本科
工作经历：
2022-2025 | 虚构运动公司 | 店长
工作职责：管理五人团队，负责排班和库存。
业绩：月度销售80万元。
离职原因：测试离职原因不导出
2018-2022 | 虚构零售公司 | 主管
工作职责：负责培训和客户服务。
业绩：培训10名同事。
期望薪资：测试薪资不导出
'''
def fictional_template(path):
 doc=Document()
 doc.sections[0].header.paragraphs[0].text='TEST TEMPLATE - FICTIONAL'
 doc.add_paragraph('Fictional candidate report')
 doc.add_paragraph('候选人姓名')
 table=doc.add_table(rows=10,cols=2)
 for row,label in zip(table.rows,['年龄','性别','城市','出生年月','手机号','邮箱','身高体重','婚育','当前薪资','期望薪资']):
  row.cells[0].text=label
 doc.add_paragraph('教育背景')
 doc.add_table(rows=1,cols=4)
 doc.add_paragraph('工作经历概况')
 table=doc.add_table(rows=2,cols=3)
 for cell,label in zip(table.rows[0].cells,['任职时间','公司','职位']):cell.text=label
 doc.add_paragraph('工作经历')
 for text in ['时间 公司','职位','工作职责：','虚构职责','业绩：','虚构业绩']:doc.add_paragraph(text)
 doc.save(path)

class V2Tests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
  self.template_path=Path(self.tmp.name)/'fictional-template.docx'
  fictional_template(self.template_path)
  # Replace only the test's default template dependency; production code is unchanged.
  renderer=patch('store_v2.render_profile',side_effect=lambda template,profile:render_profile(self.template_path if isinstance(template,Path) else template,profile))
  renderer.start();self.addCleanup(renderer.stop)
  self.s=V2Store(Path(self.tmp.name)/'db.sqlite3')
  self.j=self.s.action('job',{'title':'虚构店长','jd':'门店管理'})['id']
  self.i=self.s.v2('candidate',{'job_id':self.j,'name':'测试候选人A','resume':RESUME})['id']
 def c(self):return self.s.candidate(self.i)[1]
 def action(self,a,**kw):return self.s.v2(a,dict(id=self.i,revision=self.c().get('revision',0),**kw))
 def test_extraction_and_missing(self):
  p=self.c()['structured'];self.assertEqual(len(p['experience']),2);self.assertEqual(len(p['education']),1)
  self.assertFalse(any(x['path']=='personal.城市' for x in missing(p)))
  self.assertTrue(any(x['path']=='personal.身高' for x in missing(p)))
 def test_revision_history_and_archive(self):
  self.action('profile',profile=self.c()['structured'])
  with self.assertRaises(ValueError):self.s.v2('profile',{'id':self.i,'revision':0,'profile':self.c()['structured']})
  self.assertEqual(len(self.s.v2('history',{'id':self.i})['history']),1)
  with self.assertRaises(ValueError):self.action('decision',decision='推荐')
  self.action('debrief',objective='电话确认团队5人',complete=True)
  with self.assertRaises(ValueError):self.action('decision',decision='不推荐',reason='')
  self.action('decision',decision='不推荐',reason='此次岗位地点不适合')
  self.assertTrue(self.c()['archived'])
  with self.assertRaises(ValueError):self.action('profile',profile=self.c()['structured'])
  self.action('reopen');self.assertFalse(self.c()['archived'])
 def test_word_and_privacy(self):
  self.action('debrief',objective='虚构电话记录',complete=True);self.action('decision',decision='推荐')
  blob=self.s.word(self.i);d=Document(io.BytesIO(blob))
  text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
  self.assertIn('虚构运动公司',text);self.assertIn('虚构零售公司',text)
  self.assertNotIn('测试薪资不导出',text);self.assertNotIn('测试离职原因不导出',text);self.assertNotIn('候选人姓名',text)
  self.assertEqual(len(d.tables[2].rows),3);self.assertEqual(len(d.inline_shapes),0)
  self.assertIn('TEST TEMPLATE - FICTIONAL',d.sections[0].header.paragraphs[0].text)
 def test_schedule_cancel(self):
  self.action('schedule',time=(datetime.now(timezone.utc)+timedelta(minutes=9)).isoformat());self.s.tick()
  self.assertIn('T-10',self.c()['reminders']);self.assertTrue(self.c()['prep']['questions'])
  self.action('cancel');self.s.tick();self.assertEqual(self.c()['reminders'],[])
 def test_templates_and_multiple_jobs(self):
  blob=self.template_path.read_bytes()
  tid=self.s.v2('template-upload',{'name':'阿迪模板测试v2','data':base64.b64encode(blob).decode()})['id']
  self.s.v2('job-template',{'job_id':self.j,'template_id':tid})
  self.action('debrief',objective='测试',complete=True);self.action('decision',decision='推荐')
  self.assertTrue(self.s.word(self.i).startswith(b'PK'))
  j2=self.s.action('job',{'title':'虚构第二岗位','jd':'培训'})['id']
  r=self.s.v2('link-job',{'id':self.i,'job_id':j2})
  _,other=self.s.candidate(r['id']);self.assertEqual(other['person_id'],self.c()['person_id']);self.assertNotIn('decision',other)
  with self.assertRaises(ValueError):self.s.v2('link-job',{'id':self.i,'job_id':j2})
 def test_ai_literal_evidence(self):
  from unittest.mock import patch
  import ai_v2
  p=self.c()['structured'];p['personal']['身高']='编造身高'
  with patch('ai_v2.call',return_value=p):
   result=ai_v2.parse_resume(RESUME)
  self.assertEqual(result['personal']['身高'],'');self.assertEqual(result['personal']['城市'],'上海')
 def test_backup_restore(self):
  backup=self.s.backup();self.action('debrief',objective='later')
  self.s.restore({'data':base64.b64encode(backup).decode(),'confirmed':True})
  self.assertNotIn('debrief',self.c());self.assertTrue(list((Path(self.tmp.name)/'backups').glob('*.sqlite3')))
  with self.assertRaises(ValueError):self.s.restore({'data':base64.b64encode(b'invalid').decode(),'confirmed':True})
if __name__=='__main__':unittest.main()
